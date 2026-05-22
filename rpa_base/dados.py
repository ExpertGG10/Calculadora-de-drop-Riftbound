import csv
import os
import unicodedata
from copy import copy

from openpyxl import Workbook, load_workbook

from .excecoes import ProblemaEsperadoError


def _normalizar_chave(chave):
    texto = unicodedata.normalize("NFD", str(chave or ""))
    texto_sem_acento = "".join(caractere for caractere in texto if unicodedata.category(caractere) != "Mn")
    return texto_sem_acento.strip().lower()


def _obter_valor_codigo(registro):
    for chave, valor in registro.items():
        if _normalizar_chave(chave) == "codigo":
            return valor
    return None


def _obter_extensao(caminho_arquivo):
    return os.path.splitext(caminho_arquivo)[1].lower()


def _linha_esta_completa(valores_linha):
    indices_preenchidos = [indice for indice, valor in enumerate(valores_linha) if valor is not None and str(valor).strip() != ""]
    if not indices_preenchidos:
        return False
    inicio_intervalo = indices_preenchidos[0]
    fim_intervalo = indices_preenchidos[-1]
    for valor in valores_linha[inicio_intervalo : fim_intervalo + 1]:
        if valor is None or str(valor).strip() == "":
            return False
    return True


def _obter_intervalo_colunas_cabecalho(linha_cabecalho):
    indices_preenchidos = [indice for indice, valor in enumerate(linha_cabecalho) if valor is not None and str(valor).strip() != ""]
    if not indices_preenchidos:
        raise ValueError("Nao foi possivel identificar colunas validas para o cabecalho.")
    return indices_preenchidos[0], indices_preenchidos[-1]


def _encontrar_indice_cabecalho(matriz):
    for indice, linha in enumerate(matriz):
        if _linha_esta_completa(linha):
            return indice
    raise ValueError("Nao foi encontrada uma linha de cabecalho completa.")


def _resolver_indice_cabecalho(matriz, linha_cabecalho=None):
    if linha_cabecalho is None:
        return _encontrar_indice_cabecalho(matriz)

    indice_cabecalho = linha_cabecalho - 1
    if indice_cabecalho < 0 or indice_cabecalho >= len(matriz):
        raise ValueError("A linha de cabecalho informada esta fora do intervalo da planilha.")

    return indice_cabecalho


def _normalizar_valor(valor):
    if valor is None:
        return ""
    return valor


def _ler_tabela_texto(caminho_arquivo, delimitador=None, encoding="utf-8-sig", linha_cabecalho=None):
    with open(caminho_arquivo, "r", encoding=encoding, newline="") as arquivo:
        if delimitador is None:
            amostra = arquivo.read(2048)
            arquivo.seek(0)
            delimitador_detectado = csv.Sniffer().sniff(amostra, delimiters=",;\t|").delimiter
        else:
            delimitador_detectado = delimitador

        leitor = csv.reader(arquivo, delimiter=delimitador_detectado)
        matriz = [linha for linha in leitor]

    indice_cabecalho = _resolver_indice_cabecalho(matriz, linha_cabecalho=linha_cabecalho)
    linha_cabecalho = matriz[indice_cabecalho]
    inicio_coluna, fim_coluna = _obter_intervalo_colunas_cabecalho(linha_cabecalho)
    cabecalho = [str(coluna).strip() for coluna in linha_cabecalho[inicio_coluna : fim_coluna + 1]]
    linhas_dados = matriz[indice_cabecalho + 1 :]

    registros = []
    for linha in linhas_dados:
        linha_recortada = linha[inicio_coluna : fim_coluna + 1]
        linha_ajustada = list(linha_recortada) + [""] * (len(cabecalho) - len(linha_recortada))
        linha_ajustada = linha_ajustada[: len(cabecalho)]
        registros.append({cabecalho[indice]: _normalizar_valor(linha_ajustada[indice]) for indice in range(len(cabecalho))})

    return {
        "cabecalho": cabecalho,
        "registros": registros,
        "metadados": {
            "formato_origem": _obter_extensao(caminho_arquivo),
            "indice_cabecalho": indice_cabecalho,
            "linhas_antes_cabecalho": matriz[:indice_cabecalho],
            "delimitador": delimitador_detectado,
            "encoding": encoding,
        },
    }


def _capturar_estilo_celula(celula):
    return {
        "font": copy(celula.font),
        "fill": copy(celula.fill),
        "border": copy(celula.border),
        "alignment": copy(celula.alignment),
        "number_format": celula.number_format,
        "protection": copy(celula.protection),
    }


def _aplicar_estilo_celula(celula, estilo):
    if not estilo:
        return
    celula.font = copy(estilo["font"])
    celula.fill = copy(estilo["fill"])
    celula.border = copy(estilo["border"])
    celula.alignment = copy(estilo["alignment"])
    celula.number_format = estilo["number_format"]
    celula.protection = copy(estilo["protection"])


def _ler_tabela_xlsx(caminho_arquivo, nome_aba=None, linha_cabecalho=None):
    workbook = load_workbook(caminho_arquivo)
    planilha = workbook[nome_aba] if nome_aba else workbook.active

    matriz = [[celula.value for celula in linha] for linha in planilha.iter_rows()]
    indice_cabecalho = _resolver_indice_cabecalho(matriz, linha_cabecalho=linha_cabecalho)

    linha_cabecalho = matriz[indice_cabecalho]
    inicio_coluna, fim_coluna = _obter_intervalo_colunas_cabecalho(linha_cabecalho)
    cabecalho = [str(coluna).strip() for coluna in linha_cabecalho[inicio_coluna : fim_coluna + 1]]
    linhas_dados = matriz[indice_cabecalho + 1 :]

    registros = []
    for linha in linhas_dados:
        linha_recortada = linha[inicio_coluna : fim_coluna + 1]
        linha_ajustada = list(linha_recortada) + [""] * (len(cabecalho) - len(linha_recortada))
        linha_ajustada = linha_ajustada[: len(cabecalho)]
        registros.append({cabecalho[indice]: _normalizar_valor(linha_ajustada[indice]) for indice in range(len(cabecalho))})

    estilos_linhas_antes = []
    for linha_excel in range(1, indice_cabecalho + 1):
        linha_estilo = []
        for coluna_excel in range(inicio_coluna + 1, fim_coluna + 2):
            linha_estilo.append(_capturar_estilo_celula(planilha.cell(row=linha_excel, column=coluna_excel)))
        estilos_linhas_antes.append(linha_estilo)

    estilos_cabecalho = []
    linha_cabecalho_excel = indice_cabecalho + 1
    for coluna_excel in range(inicio_coluna + 1, fim_coluna + 2):
        estilos_cabecalho.append(_capturar_estilo_celula(planilha.cell(row=linha_cabecalho_excel, column=coluna_excel)))

    larguras_colunas = {}
    for coluna_excel in range(inicio_coluna + 1, fim_coluna + 2):
        letra_coluna = planilha.cell(row=linha_cabecalho_excel, column=coluna_excel).column_letter
        largura = planilha.column_dimensions[letra_coluna].width
        if largura is not None:
            larguras_colunas[letra_coluna] = largura

    alturas_linhas = {}
    for linha_excel in range(1, indice_cabecalho + 2):
        altura = planilha.row_dimensions[linha_excel].height
        if altura is not None:
            alturas_linhas[linha_excel] = altura

    return {
        "cabecalho": cabecalho,
        "registros": registros,
        "metadados": {
            "formato_origem": ".xlsx",
            "nome_aba": planilha.title,
            "indice_cabecalho": indice_cabecalho,
            "linhas_antes_cabecalho": matriz[:indice_cabecalho],
            "estilos_linhas_antes_cabecalho": estilos_linhas_antes,
            "estilos_cabecalho": estilos_cabecalho,
            "larguras_colunas": larguras_colunas,
            "alturas_linhas": alturas_linhas,
        },
    }


def planilha_para_dicionario(caminho_arquivo, delimitador=None, encoding="utf-8-sig", nome_aba=None, linha_cabecalho=None):
    extensao = _obter_extensao(caminho_arquivo)

    if extensao == ".xlsx":
        return _ler_tabela_xlsx(caminho_arquivo, nome_aba=nome_aba, linha_cabecalho=linha_cabecalho)
    if extensao in [".csv", ".tsv", ".txt"]:
        delimitador_arquivo = "\t" if extensao == ".tsv" and delimitador is None else delimitador
        return _ler_tabela_texto(
            caminho_arquivo,
            delimitador=delimitador_arquivo,
            encoding=encoding,
            linha_cabecalho=linha_cabecalho,
        )

    raise ValueError(f"Formato nao suportado: {extensao}")


def dicionario_para_planilha(dados, caminho_saida, delimitador=None, encoding="utf-8-sig"):
    cabecalho = dados.get("cabecalho", [])
    registros = dados.get("registros", [])
    metadados = dados.get("metadados", {})

    if not cabecalho:
        raise ValueError("O dicionario precisa conter um cabecalho nao vazio.")

    extensao = _obter_extensao(caminho_saida)

    if extensao == ".xlsx":
        workbook = Workbook()
        nome_aba = metadados.get("nome_aba", "Planilha1")
        planilha = workbook.active
        planilha.title = nome_aba

        linhas_antes_cabecalho = metadados.get("linhas_antes_cabecalho", [])
        estilos_linhas_antes = metadados.get("estilos_linhas_antes_cabecalho", [])
        estilos_cabecalho = metadados.get("estilos_cabecalho", [])

        linha_atual = 1
        for indice_linha, linha in enumerate(linhas_antes_cabecalho):
            for indice_coluna, valor in enumerate(linha, start=1):
                celula = planilha.cell(row=linha_atual, column=indice_coluna, value=valor)
                if indice_linha < len(estilos_linhas_antes) and indice_coluna - 1 < len(estilos_linhas_antes[indice_linha]):
                    _aplicar_estilo_celula(celula, estilos_linhas_antes[indice_linha][indice_coluna - 1])
            linha_atual += 1

        for indice_coluna, coluna in enumerate(cabecalho, start=1):
            celula = planilha.cell(row=linha_atual, column=indice_coluna, value=coluna)
            if indice_coluna - 1 < len(estilos_cabecalho):
                _aplicar_estilo_celula(celula, estilos_cabecalho[indice_coluna - 1])

        for registro in registros:
            linha_atual += 1
            for indice_coluna, coluna in enumerate(cabecalho, start=1):
                planilha.cell(row=linha_atual, column=indice_coluna, value=registro.get(coluna, ""))

        for letra_coluna, largura in metadados.get("larguras_colunas", {}).items():
            planilha.column_dimensions[letra_coluna].width = largura

        for numero_linha, altura in metadados.get("alturas_linhas", {}).items():
            planilha.row_dimensions[numero_linha].height = altura

        workbook.save(caminho_saida)
        return caminho_saida

    if extensao in [".csv", ".tsv", ".txt"]:
        if delimitador is None:
            delimitador = "\t" if extensao == ".tsv" else metadados.get("delimitador", ",")

        with open(caminho_saida, "w", encoding=encoding, newline="") as arquivo:
            escritor = csv.writer(arquivo, delimiter=delimitador)
            for linha in metadados.get("linhas_antes_cabecalho", []):
                escritor.writerow(linha)
            escritor.writerow(cabecalho)
            for registro in registros:
                escritor.writerow([registro.get(coluna, "") for coluna in cabecalho])

        return caminho_saida

    raise ValueError(f"Formato de saida nao suportado: {extensao}")




def obter_id_empresa(registro, indice):
    valor = _obter_valor_codigo(registro)
    if valor is not None and str(valor).strip() != "":
        return str(valor).strip()
    return f"linha_{indice}"


def obter_codigo(registro, indice, fluxo_codigo_invalido=None):
    valor = _obter_valor_codigo(registro)
    if valor is not None and str(valor).strip() != "":
        codigo = str(valor).strip()
        if "/" in codigo:
            codigos = [parte.strip() for parte in codigo.split("/") if parte.strip()]
            if codigos:
                return codigos
        return codigo
    raise ProblemaEsperadoError("Código não encontrado", fluxo=fluxo_codigo_invalido, indice=indice)


__all__ = ["planilha_para_dicionario", "dicionario_para_planilha", "obter_id_empresa", "obter_codigo"]
